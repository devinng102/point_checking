"""Check CSV/XLSX point lists against MongoDB and append Excel results."""

import argparse
import csv
import json
import sys
import tempfile
from collections import Counter
from collections.abc import Sequence
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.cell.rich_text import CellRichText
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.fills import Fill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet
from pymongo import MongoClient
from pymongo.errors import PyMongoError

from building_config import SCHEMA_DATABASE_ALIASES
from checking_logic import (
    BUILDING_COLUMN,
    EQUIPMENT_COLUMN,
    POINT_COLUMN,
    CheckingError,
    CheckingResult,
    Document,
    PointChecker,
    Status,
    required_text,
)

DEFAULT_INPUT = Path(__file__).with_name("point_name_16-9-2026_v1.csv")
REQUIRED_COLUMNS = (BUILDING_COLUMN, EQUIPMENT_COLUMN, POINT_COLUMN)
REMARK_COLUMN = "Remark"
REVIEW_COLUMN = "Review"
RemarkKey = tuple[str, str]


@dataclass(frozen=True)
class ManualReview:
    text: str = ""
    fill: Fill = field(default_factory=PatternFill)


RESULT_COLUMNS = (
    "Appendix_b",
    "enm_semantic_schema",
    "Test result",
    "Derived Function Code",
    "Derived Point Base Name",
    "Base Function Code",
    "Device ID",
    "Appendix Match Method",
    "Checking Reasons",
    "Candidate Mappings",
)
RESULT_COLORS: dict[Status, str] = {
    "matched": "D5E8D4",
    "unmatched": "F8CECC",
    "review": "FFF2CC",
    "multiple mappings": "DAE8FC",
}


class MongoSchemaSource:
    def __init__(self, client: MongoClient[dict[str, Any]]) -> None:
        self.client = client
        self._cache: dict[str, Sequence[Document] | CheckingError] = {}

    def _read(self, database: str, collection: str) -> list[Document]:
        try:
            db = self.client[database]
            if collection not in db.list_collection_names(
                filter={"name": collection}
            ):
                raise CheckingError(
                    f"MongoDB collection {database}.{collection} does not exist."
                )
            return list(db[collection].find())
        except PyMongoError as exc:
            raise CheckingError(
                f"MongoDB read failed for {database}.{collection} "
                f"({type(exc).__name__}): {exc}"
            ) from exc

    def get_appendix(self) -> list[Document]:
        return self._read("MASTER", "appendix_b")

    def get_schema(self, building: str) -> Sequence[Document]:
        building = SCHEMA_DATABASE_ALIASES.get(building, building)
        if building not in self._cache:
            try:
                self._cache[building] = self._read(
                    building, "enm_semantic_schema"
                )
            except CheckingError as exc:
                self._cache[building] = exc
        cached = self._cache[building]
        if isinstance(cached, CheckingError):
            raise cached
        return cached


def write_text(cell: Cell, value: str) -> None:
    if len(value) > 32_767 or ILLEGAL_CHARACTERS_RE.search(value):
        raise ValueError(
            f"{cell.parent.title}!{cell.coordinate}: text cannot be represented "
            "in Excel without changing its value."
        )
    cell.value = value
    cell.data_type = "s"


def read_workbook(path: Path) -> tuple[Workbook, Workbook | None]:
    if path.suffix.lower() == ".csv":
        workbook = Workbook()
        sheet = workbook.active
        if not isinstance(sheet, Worksheet):
            raise ValueError("Could not create a worksheet.")
        sheet.title = "Point list"
        try:
            with path.open(encoding="utf-8-sig", newline="") as source:
                for row_number, row in enumerate(csv.reader(source), start=1):
                    for column_number, value in enumerate(row, start=1):
                        write_text(sheet.cell(row_number, column_number), value)
        except (OSError, UnicodeError, csv.Error, ValueError):
            workbook.close()
            raise
        return workbook, None
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("Input must be a .csv, .xlsx or .xlsm file.")
    keep_vba = path.suffix.lower() == ".xlsm"
    workbook = load_workbook(path, keep_vba=keep_vba, rich_text=True)
    try:
        cached = load_workbook(path, data_only=True, read_only=True)
    except (OSError, ValueError, BadZipFile, InvalidFileException):
        workbook.close()
        raise
    return workbook, cached


def input_columns(
    sheet: Worksheet, required: Sequence[str] = REQUIRED_COLUMNS
) -> dict[str, int] | None:
    columns: dict[str, int] = {}
    for cell in sheet[1]:
        if cell.value in required:
            if cell.value in columns:
                raise ValueError(
                    f"{sheet.title}: duplicate required header {cell.value!r}."
                )
            columns[cell.value] = cell.column
    return columns if all(name in columns for name in required) else None


def input_value(cell: Cell, cached: Workbook | None) -> Any:
    if cell.data_type != "f":
        return cell.value
    value = (
        cached[cell.parent.title].cell(cell.row, cell.column).value
        if cached is not None else None
    )
    if value is None:
        raise CheckingError(
            f"{cell.parent.title}!{cell.coordinate} is a formula without "
            "a cached value; recalculate the source in Excel."
        )
    return value


def manual_text(cell: Cell, name: str) -> str:
    if cell.value in (None, ""):
        return ""
    if cell.data_type == "f" or not isinstance(cell.value, (str, CellRichText)):
        raise ValueError(
            f"{cell.parent.title}!{cell.coordinate}: {name} must be "
            "literal text, not a formula or another data type."
        )
    return str(cell.value)


def manual_review(cell: Cell) -> ManualReview:
    return ManualReview(manual_text(cell, REVIEW_COLUMN), copy(cell.fill))


def read_manual_annotations(
    path: Path,
) -> tuple[dict[RemarkKey, str], dict[RemarkKey, ManualReview]]:
    workbook, cached = read_workbook(path)
    remarks: dict[RemarkKey, str] = {}
    reviews: dict[RemarkKey, ManualReview] = {}
    found_sheet = False
    try:
        for sheet in workbook.worksheets:
            manual_columns = [
                name for name in (REMARK_COLUMN, REVIEW_COLUMN)
                if any(cell.value == name for cell in sheet[1])
            ]
            if not manual_columns:
                continue
            columns = input_columns(
                sheet, (BUILDING_COLUMN, POINT_COLUMN, *manual_columns)
            )
            if columns is None:
                raise ValueError(
                    f"{sheet.title}: {' / '.join(manual_columns)} requires "
                    f"{BUILDING_COLUMN} and "
                    f"{POINT_COLUMN} headers."
                )
            found_sheet = True
            for row in range(2, sheet.max_row + 1):
                remark = (
                    manual_text(sheet.cell(row, columns[REMARK_COLUMN]), REMARK_COLUMN)
                    if REMARK_COLUMN in columns else ""
                )
                review = (
                    manual_review(sheet.cell(row, columns[REVIEW_COLUMN]))
                    if REVIEW_COLUMN in columns else ManualReview()
                )
                if not remark and review == ManualReview():
                    continue
                key = (
                    required_text(
                        input_value(sheet.cell(row, columns[BUILDING_COLUMN]), cached),
                        f"{sheet.title}!row {row}: {BUILDING_COLUMN}",
                    ),
                    required_text(
                        input_value(sheet.cell(row, columns[POINT_COLUMN]), cached),
                        f"{sheet.title}!row {row}: {POINT_COLUMN}",
                    ),
                )
                if remark and key in remarks and remarks[key] != remark:
                    raise ValueError(
                        f"Conflicting Remarks for {key!r} in {path}: "
                        "resolve the duplicate notes before importing."
                    )
                if remark:
                    remarks[key] = remark
                if review != ManualReview():
                    if key in reviews and reviews[key] != review:
                        raise ValueError(
                            f"Conflicting Reviews for {key!r} in {path}: "
                            "resolve duplicate review text/backgrounds before importing."
                        )
                    reviews[key] = review
        if not found_sheet:
            raise ValueError(
                f"No worksheet in {path} contains a Remark or Review column with "
                f"{BUILDING_COLUMN} and {POINT_COLUMN}."
            )
        return remarks, reviews
    finally:
        if cached is not None:
            cached.close()
        workbook.close()


def read_remarks(path: Path) -> dict[RemarkKey, str]:
    remarks, _ = read_manual_annotations(path)
    return remarks


def result_values(result: CheckingResult) -> tuple[str, ...]:
    def boolean(value: bool | None) -> str:
        return "" if value is None else ("T" if value else "F")

    return (
        boolean(result.appendix),
        boolean(result.schema),
        result.status,
        result.function_code,
        result.point_base_name,
        result.base_function_code,
        result.device_id,
        result.match_method,
        "\n".join([*result.reasons, *result.errors]),
        json.dumps(result.candidate_results, ensure_ascii=False)
        if result.candidate_results else "",
    )


def check_workbook(
    workbook: Workbook,
    checker: PointChecker,
    cached: Workbook | None = None,
    remarks: dict[RemarkKey, str] | None = None,
    reviews: dict[RemarkKey, ManualReview] | None = None,
) -> Counter[Status]:
    summaries: Counter[Status] = Counter()
    checked_sheets = 0
    matched_remarks: set[RemarkKey] = set()
    matched_reviews: set[RemarkKey] = set()
    for sheet in workbook.worksheets:
        columns = input_columns(sheet)
        if columns is None:
            continue
        checked_sheets += 1
        source_width = sheet.max_column
        source_height = sheet.max_row
        remark_columns = [
            cell.column for cell in sheet[1] if cell.value == REMARK_COLUMN
        ]
        if len(remark_columns) > 1:
            raise ValueError(f"{sheet.title}: duplicate Remark header.")
        review_columns = [
            cell.column for cell in sheet[1] if cell.value == REVIEW_COLUMN
        ]
        if len(review_columns) > 1:
            raise ValueError(f"{sheet.title}: duplicate Review header.")
        used_headers = {str(cell.value) for cell in sheet[1] if cell.value is not None}
        for offset, name in enumerate(RESULT_COLUMNS, start=1):
            header = name
            suffix = 2
            while header in used_headers:
                header = f"{name} (checking {suffix})"
                suffix += 1
            used_headers.add(header)
            cell = sheet.cell(1, source_width + offset)
            write_text(cell, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E0F2F1")
            sheet.column_dimensions[get_column_letter(cell.column)].width = (
                55 if name in ("Checking Reasons", "Candidate Mappings") else 26
            )

        remark_column = (
            remark_columns[0] if remark_columns
            else source_width + len(RESULT_COLUMNS) + 1
        )
        if not remark_columns:
            cell = sheet.cell(1, remark_column)
            write_text(cell, REMARK_COLUMN)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E0F2F1")
            sheet.column_dimensions[get_column_letter(remark_column)].width = 55

        review_column = (
            review_columns[0] if review_columns
            else source_width + len(RESULT_COLUMNS) + (1 if remark_columns else 2)
        )
        if not review_columns:
            cell = sheet.cell(1, review_column)
            write_text(cell, REVIEW_COLUMN)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E0F2F1")
            sheet.column_dimensions[get_column_letter(review_column)].width = 26

        for row_number in range(2, source_height + 1):
            if all(
                sheet.cell(row_number, column).value in (None, "")
                for column in range(1, source_width + 1)
            ):
                continue
            values: dict[str, Any] = {}
            formula_errors: list[str] = []
            for name, column in columns.items():
                cell = sheet.cell(row_number, column)
                try:
                    values[name] = input_value(cell, cached)
                except CheckingError as exc:
                    values[name] = None
                    formula_errors.append(f"{name}: {exc}")
            result = (
                CheckingResult(errors=formula_errors)
                if formula_errors else checker.check(values)
            )
            summaries[result.status] += 1
            for offset, value in enumerate(result_values(result), start=1):
                cell = sheet.cell(row_number, source_width + offset)
                write_text(cell, value)
                cell.fill = PatternFill(
                    "solid", fgColor=RESULT_COLORS[result.status]
                )
            remark_cell = sheet.cell(row_number, remark_column)
            review_cell = sheet.cell(row_number, review_column)
            key_values = (values[BUILDING_COLUMN], values[POINT_COLUMN])
            if (remarks is not None or reviews is not None) and all(
                isinstance(value, str) and value.strip() for value in key_values
            ):
                key = (
                    required_text(key_values[0], BUILDING_COLUMN),
                    required_text(key_values[1], POINT_COLUMN),
                )
                if remarks is not None and key in remarks:
                    if (
                        remark_cell.value not in (None, "", remarks[key])
                        or remark_cell.data_type == "f"
                    ):
                        raise ValueError(
                            f"{sheet.title}!{remark_cell.coordinate}: existing "
                            f"Remark conflicts with imported Remark for {key!r}; "
                            "resolve the notes before exporting."
                        )
                    if remark_cell.value in (None, ""):
                        write_text(remark_cell, remarks[key])
                    matched_remarks.add(key)
                if reviews is not None and key in reviews:
                    current_review = manual_review(review_cell)
                    imported_review = reviews[key]
                    if current_review not in (ManualReview(), imported_review):
                        raise ValueError(
                            f"{sheet.title}!{review_cell.coordinate}: existing "
                            f"Review conflicts with imported Review for {key!r}; "
                            "resolve text/backgrounds before exporting."
                        )
                    if current_review == ManualReview():
                        write_text(review_cell, imported_review.text)
                        review_cell.fill = copy(imported_review.fill)
                    matched_reviews.add(key)
            if not review_columns and review_cell.value is None:
                write_text(review_cell, "")
            if not remark_columns:
                if remark_cell.value is None:
                    write_text(remark_cell, "")
                remark_cell.fill = PatternFill(
                    "solid", fgColor=RESULT_COLORS[result.status]
                )
            for error in result.errors:
                print(
                    f"{sheet.title}!row {row_number}: {error}",
                    file=sys.stderr,
                )
    if not checked_sheets:
        raise ValueError(
            "No worksheet contains all required headers: "
            + ", ".join(REQUIRED_COLUMNS)
            + ". Headers must be on the first row."
        )
    if remarks is not None:
        print(
            f"Remarks: matched {len(matched_remarks)} of {len(remarks)} "
            "annotated Building / Point Name keys."
        )
        for key in sorted(remarks.keys() - matched_remarks):
            print(
                f"Warning: Remark for {key!r} has no matching input row; "
                "it remains only in the previous remarks file.",
                file=sys.stderr,
            )
    if reviews is not None:
        print(
            f"Reviews: matched {len(matched_reviews)} of {len(reviews)} "
            "reviewed Building / Point Name keys."
        )
        for key in sorted(reviews.keys() - matched_reviews):
            print(
                f"Warning: Review for {key!r} has no matching input row; "
                "it remains only in the previous annotations file.",
                file=sys.stderr,
            )
    return summaries


def save_workbook(workbook: Workbook, output: Path) -> None:
    if output.exists():
        raise FileExistsError(f"Output already exists: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=output.parent, prefix=f".{output.stem}-",
            suffix=output.suffix, delete=False,
        ) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        output.hardlink_to(temporary)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Check a CSV/XLSX point list against MongoDB and append colored "
            "Excel results without changing original columns."
        )
    )
    parser.add_argument(
        "input_file", nargs="?", type=Path, default=DEFAULT_INPUT,
        help="input CSV/XLSX/XLSM; headers must be on row 1",
    )
    parser.add_argument(
        "-o", "--output", type=Path,
        help="new output workbook; existing files are never overwritten",
    )
    parser.add_argument(
        "--remarks-from", type=Path,
        help=(
            "previous CSV/XLSX/XLSM with manual Remark and/or Review; preserve "
            "Review text/background by exact Building & Block Code and full Point Name"
        ),
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source_path: Path = args.input_file
    suffix = ".xlsm" if source_path.suffix.lower() == ".xlsm" else ".xlsx"
    output = args.output or source_path.with_name(
        f"{source_path.stem}_checked{suffix}"
    )
    workbook: Workbook | None = None
    cached: Workbook | None = None
    client: MongoClient[dict[str, Any]] | None = None
    try:
        if output.resolve() == source_path.resolve():
            raise ValueError("Output must not overwrite the input file.")
        if output.suffix.lower() != suffix:
            raise ValueError(f"Output must use the {suffix} extension.")
        if output.exists():
            raise FileExistsError(f"Output already exists: {output}")
        if (
            args.remarks_from is not None
            and output.resolve() == args.remarks_from.resolve()
        ):
            raise ValueError("Output must not overwrite the remarks source.")
        remarks, reviews = (
            read_manual_annotations(args.remarks_from)
            if args.remarks_from is not None else (None, None)
        )
        workbook, cached = read_workbook(source_path)
        if not any(input_columns(sheet) for sheet in workbook.worksheets):
            raise ValueError("Input has no worksheet with the required headers.")

        from mongodb_connect import DB

        client = DB
        client.admin.command("ping")
        data = MongoSchemaSource(client)
        checker = PointChecker(data.get_appendix(), data)
        summaries = check_workbook(workbook, checker, cached, remarks, reviews)
        save_workbook(workbook, output)
    except (
        OSError, ValueError, UnicodeError, csv.Error, BadZipFile,
        InvalidFileException, PyMongoError,
    ) as exc:
        print(
            f"Checking export failed ({type(exc).__name__}): {exc}",
            file=sys.stderr,
        )
        return 1
    finally:
        if cached is not None:
            cached.close()
        if workbook is not None:
            workbook.close()
        if client is not None:
            client.close()

    print(f"Exported {sum(summaries.values())} checked rows to {output.resolve()}")
    for status in RESULT_COLORS:
        print(f"{status}: {summaries[status]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
