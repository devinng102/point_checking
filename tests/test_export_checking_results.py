import argparse
import contextlib
import csv
import io
import json
import sys
import tempfile
import unittest
from copy import copy
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from pymongo.errors import OperationFailure

from checking_logic import (
    BUILDING_COLUMN, EQUIPMENT_COLUMN, POINT_COLUMN, CheckingError, PointChecker,
)
from export_checking_results import (
    RESULT_COLORS,
    RESULT_COLUMNS,
    REMARK_COLUMN,
    MongoSchemaSource,
    check_workbook,
    main,
    read_workbook,
    read_remarks,
    parse_args,
    save_workbook,
    write_text,
)


HEADERS = [
    BUILDING_COLUMN, EQUIPMENT_COLUMN, POINT_COLUMN,
    "Function or Status Code", "point name (without function code)", "Comment",
]
APPENDIX = [{
    "Function or Status Code": "S",
    "Equipment Code": "CHR/ACC/WCC",
    "System": "HVAC/W",
    "Equipment": "Chiller",
    "iBMS Parameters": "Off/On Status",
}]
SCHEMA = [{
    "HVAC/W": {"Chiller": {"APB-ACC-04": {
        "Off/On Status": {"point_name": "APB-ACC-04-S"}
    }}}
}]


class FakeCollection:
    def __init__(self, documents):
        self.documents = documents
        self.reads = 0

    def find(self):
        self.reads += 1
        return iter(self.documents)


class FakeDatabase:
    def __init__(self, collections, error=None):
        self.collections = collections
        self.error = error

    def list_collection_names(self, filter):
        if self.error:
            raise self.error
        return list(self.collections)

    def __getitem__(self, name):
        return self.collections[name]


class FakeClient:
    def __init__(self, buildings=None):
        self.databases = {
            "MASTER": FakeDatabase({"appendix_b": FakeCollection(APPENDIX)}),
            "APB": FakeDatabase({"enm_semantic_schema": FakeCollection(SCHEMA)}),
        }
        if buildings:
            self.databases.update(buildings)
        self.closed = False
        self.admin = SimpleNamespace(command=lambda name: None)

    def __getitem__(self, name):
        return self.databases.get(name, FakeDatabase({}))

    def close(self):
        self.closed = True


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary.name)
        self.client = FakeClient()
        data = MongoSchemaSource(self.client)
        self.checker = PointChecker(data.get_appendix(), data)

    def tearDown(self):
        self.temporary.cleanup()

    def workbook(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(HEADERS)
        sheet.append(["APB", "ACC", "APB-ACC-04-S", "WRONG", "WRONG", "001"])
        return workbook

    def remarks_file(self, rows, headers=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [BUILDING_COLUMN, POINT_COLUMN, REMARK_COLUMN]
            if headers is None else headers
        )
        for values in rows:
            sheet.append(values)
        path = self.directory / "previous.xlsx"
        workbook.save(path)
        workbook.close()
        return path

    def test_preserves_original_values_styles_and_other_sheets(self):
        workbook = self.workbook()
        sheet = workbook.active
        sheet["C2"].fill = PatternFill("solid", fgColor="123456")
        notes = workbook.create_sheet("Notes")
        notes.append(["Untouched", "=1+2"])
        source = self.directory / "source.xlsx"
        workbook.save(source)
        workbook.close()

        workbook, cached = read_workbook(source)
        summary = check_workbook(workbook, self.checker, cached)
        output = self.directory / "checked.xlsx"
        save_workbook(workbook, output)
        workbook.close()
        cached.close()

        actual = load_workbook(output)
        original = load_workbook(source)
        for title in original.sheetnames:
            for rows in original[title].iter_rows():
                for cell in rows:
                    target = actual[title][cell.coordinate]
                    self.assertEqual(target.value, cell.value)
                    self.assertEqual(target.data_type, cell.data_type)
                    self.assertEqual(copy(target.fill), copy(cell.fill))
        self.assertEqual(summary["matched"], 1)
        self.assertEqual(actual.active["G2"].value, "T")
        self.assertEqual(actual.active["H2"].value, "T")
        self.assertEqual(actual.active["I2"].value, "matched")
        self.assertEqual(actual.active["J2"].value, "S")
        self.assertEqual(actual.active["I2"].fill.fgColor.rgb[-6:], "D5E8D4")
        actual.close()
        original.close()

    def test_csv_text_preserves_leading_zeroes_and_literal_formulas(self):
        source = self.directory / "source.csv"
        with source.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(HEADERS)
            writer.writerow([
                "APB", "ACC", "APB-ACC-04-S", "WRONG", "WRONG", "=1+2"
            ])
            writer.writerow([
                "APB", "ACC", "APB-ACC-04-XYZ", "WRONG", "WRONG", "001"
            ])
        workbook, cached = read_workbook(source)
        self.assertIsNone(cached)
        check_workbook(workbook, self.checker)
        output = self.directory / "checked.xlsx"
        save_workbook(workbook, output)
        workbook.close()
        actual = load_workbook(output)
        self.assertEqual(actual.active["F2"].value, "=1+2")
        self.assertEqual(actual.active["F2"].data_type, "s")
        self.assertEqual(actual.active["F3"].value, "001")
        self.assertEqual(actual.active["D2"].value, "WRONG")
        self.assertEqual(actual.active["E2"].value, "WRONG")
        actual.close()

    def test_all_four_result_backgrounds(self):
        records = APPENDIX + [{
            **APPENDIX[0], "Function or Status Code": "MULTI",
        }, {
            **APPENDIX[0], "Function or Status Code": "MULTI",
            "iBMS Parameters": "Other",
        }]
        data = MongoSchemaSource(self.client)
        checker = PointChecker(records, data)
        workbook = self.workbook()
        sheet = workbook.active
        sheet.append(["APB", "ACC", "APB-ACC-05-S"])
        sheet.append(["APB", "ACC", "APB-ACC-04-SWT"])
        sheet.append(["APB", "ACC", "APB-ACC-04-MULTI"])
        summary = check_workbook(workbook, checker)
        self.assertEqual(dict(summary), {
            "matched": 1, "unmatched": 1, "review": 1, "multiple mappings": 1,
        })
        for row_number, status in enumerate(RESULT_COLORS, start=2):
            self.assertEqual(sheet.cell(row_number, 9).value, status)
            for column in range(7, 7 + len(RESULT_COLUMNS)):
                self.assertEqual(
                    sheet.cell(row_number, column).fill.fgColor.rgb[-6:],
                    RESULT_COLORS[status],
                )
        workbook.close()

    def test_misc_unknown_round_trip_blank_appendix_yellow_and_preserved_source(self):
        records = [{
            "Function or Status Code": "TFALM",
            "Equipment Code": None,
            "System": "PV",
            "Equipment": "PV System",
            "iBMS Parameters": "PV System Normal / Fault Alarm",
        }]
        data = MongoSchemaSource(self.client)
        checker = PointChecker(records, data)
        workbook = self.workbook()
        sheet = workbook.active
        sheet["B2"] = "SVRRM"
        sheet["C2"] = "APB-4F-SVRRM-OUTD-TFALM"
        originals = [sheet.cell(2, c).value for c in range(1, 7)]
        summary = check_workbook(workbook, checker)
        self.assertEqual(summary["review"], 1)
        output = self.directory / "misc_checked.xlsx"
        save_workbook(workbook, output)
        workbook.close()
        actual = load_workbook(output)
        self.addCleanup(actual.close)
        sheet = actual.active
        self.assertIn(sheet["G2"].value, (None, ""))
        self.assertEqual(sheet["H2"].value, "F")
        self.assertEqual(sheet["I2"].value, "review")
        self.assertEqual([sheet.cell(2, c).value for c in range(1, 7)], originals)
        for column in range(7, 7 + len(RESULT_COLUMNS)):
            self.assertEqual(sheet.cell(2, column).fill.fgColor.rgb[-6:], "FFF2CC")
        self.assertIn("full Point Name scan", sheet["O2"].value)
        candidate = json.loads(sheet["P2"].value)[0]
        self.assertIsNone(candidate["MISC Relevance"])
        self.assertIsNone(candidate["Schema Match"])
        self.assertIn("PV", candidate["Relevance Reason"])

    def test_misc_missing_collection_exports_unknown_not_false(self):
        data = MongoSchemaSource(self.client)
        checker = PointChecker([{**APPENDIX[0], "Equipment Code": None}], data)
        workbook = self.workbook()
        self.addCleanup(workbook.close)
        workbook.active["A2"] = "MISSING"
        with contextlib.redirect_stderr(io.StringIO()) as errors:
            summary = check_workbook(workbook, checker)
        self.assertEqual(summary["review"], 1)
        self.assertEqual(workbook.active["G2"].value, "")
        self.assertEqual(workbook.active["H2"].value, "")
        self.assertIn("does not exist", errors.getvalue())

    def test_existing_result_headers_are_not_overwritten(self):
        workbook = self.workbook()
        sheet = workbook.active
        sheet["G1"] = "Appendix_b"
        sheet["G2"] = "ORIGINAL"
        check_workbook(workbook, self.checker)
        self.assertEqual(sheet["G2"].value, "ORIGINAL")
        self.assertEqual(sheet["H1"].value, "Appendix_b (checking 2)")
        self.assertEqual(sheet["H2"].value, "T")
        workbook.close()

    def test_multiple_data_sheets_and_blank_rows(self):
        workbook = self.workbook()
        second = workbook.create_sheet("Second")
        second.append(HEADERS)
        second.append([None] * len(HEADERS))
        second.append(["APB", "ACC", "APB-ACC-04-S"])
        summary = check_workbook(workbook, self.checker)
        self.assertEqual(sum(summary.values()), 2)
        self.assertIsNone(second["G2"].value)
        self.assertEqual(second["G3"].value, "T")
        workbook.close()

    def test_blank_csv_rows_are_not_checked(self):
        source = self.directory / "source.csv"
        with source.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(HEADERS)
            writer.writerow([""] * len(HEADERS))
            writer.writerow(["APB", "ACC", "APB-ACC-04-S"])
        workbook, _ = read_workbook(source)
        summary = check_workbook(workbook, self.checker)
        self.assertEqual(sum(summary.values()), 1)
        self.assertIsNone(workbook.active["G2"].value)
        self.assertEqual(workbook.active["G3"].value, "T")
        workbook.close()

    def test_input_formula_without_cache_is_yellow_and_formula_preserved(self):
        workbook = self.workbook()
        workbook.active["C2"] = '="APB-ACC-04-S"'
        with contextlib.redirect_stderr(io.StringIO()) as errors:
            summary = check_workbook(workbook, self.checker)
        self.assertEqual(summary["review"], 1)
        self.assertEqual(workbook.active["C2"].value, '="APB-ACC-04-S"')
        self.assertEqual(workbook.active["G2"].value, "")
        self.assertEqual(workbook.active["H2"].value, "")
        self.assertIn("without a cached value", errors.getvalue())
        workbook.close()

    def test_input_formula_uses_cache_without_changing_formula(self):
        workbook = self.workbook()
        workbook.active["C2"] = '="APB-ACC-04-S"'
        cached = self.workbook()
        summary = check_workbook(workbook, self.checker, cached)
        self.assertEqual(summary["matched"], 1)
        self.assertEqual(workbook.active["C2"].value, '="APB-ACC-04-S"')
        workbook.close()
        cached.close()

    def test_missing_schema_collection_is_yellow_with_blank_schema(self):
        workbook = self.workbook()
        workbook.active["A2"] = "NONEXISTENT"
        with contextlib.redirect_stderr(io.StringIO()) as errors:
            summary = check_workbook(workbook, self.checker)
        self.assertEqual(summary["review"], 1)
        self.assertEqual(workbook.active["G2"].value, "T")
        self.assertEqual(workbook.active["H2"].value, "")
        self.assertIn("does not exist", errors.getvalue())
        workbook.close()

    def test_schema_reads_and_failures_are_cached_per_building(self):
        source = MongoSchemaSource(self.client)
        source.get_schema("APB")
        source.get_schema("APB")
        collection = self.client.databases["APB"].collections["enm_semantic_schema"]
        self.assertEqual(collection.reads, 1)
        for _ in range(2):
            with self.assertRaises(CheckingError):
                source.get_schema("MISSING")
        self.assertIsInstance(source._cache["MISSING"], CheckingError)

    def test_mongo_read_failure_has_context(self):
        client = FakeClient({"APB": FakeDatabase({}, OperationFailure("denied"))})
        with self.assertRaisesRegex(CheckingError, "APB.enm_semantic_schema"):
            MongoSchemaSource(client).get_schema("APB")

    def test_missing_and_duplicate_required_headers_fail(self):
        workbook = Workbook()
        workbook.active.append(["Point Name"])
        with self.assertRaisesRegex(ValueError, "No worksheet"):
            check_workbook(workbook, self.checker)
        workbook.close()
        workbook = self.workbook()
        workbook.active["G1"] = POINT_COLUMN
        with self.assertRaisesRegex(ValueError, "duplicate required header"):
            check_workbook(workbook, self.checker)
        workbook.close()

    def test_unrepresentable_text_is_not_silently_truncated(self):
        workbook = Workbook()
        for text in ("X" * 32_768, "bad\x00value"):
            with self.subTest(text_length=len(text)):
                with self.assertRaisesRegex(ValueError, "without changing"):
                    write_text(workbook.active["A1"], text)
        workbook.close()

    def test_existing_output_is_never_overwritten(self):
        workbook = self.workbook()
        output = self.directory / "checked.xlsx"
        output.write_bytes(b"original")
        with self.assertRaises(FileExistsError):
            save_workbook(workbook, output)
        self.assertEqual(output.read_bytes(), b"original")
        self.assertEqual(list(self.directory.iterdir()), [output])
        workbook.close()

    def test_failed_save_cleans_temporary_file(self):
        workbook = self.workbook()
        output = self.directory / "checked.xlsx"
        with patch.object(workbook, "save", side_effect=OSError("disk error")):
            with self.assertRaises(OSError):
                save_workbook(workbook, output)
        self.assertFalse(output.exists())
        self.assertEqual(list(self.directory.iterdir()), [])
        workbook.close()

    def test_output_created_during_save_is_never_overwritten(self):
        workbook = self.workbook()
        output = self.directory / "checked.xlsx"
        real_save = workbook.save

        def create_competing_file(path):
            real_save(path)
            output.write_bytes(b"other writer")

        with patch.object(workbook, "save", side_effect=create_competing_file):
            with self.assertRaises(FileExistsError):
                save_workbook(workbook, output)
        self.assertEqual(output.read_bytes(), b"other writer")
        self.assertEqual(list(self.directory.iterdir()), [output])
        workbook.close()

    def test_cli_exports_and_closes_client(self):
        source = self.directory / "source.csv"
        with source.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(HEADERS)
            writer.writerow(["APB", "ACC", "APB-ACC-04-S", "WRONG", "WRONG", "001"])
        output = self.directory / "checked.xlsx"
        with (
            patch("export_checking_results.parse_args", return_value=argparse.Namespace(
                input_file=source, output=output, remarks_from=None,
            )),
            patch.dict(sys.modules, {"mongodb_connect": SimpleNamespace(DB=self.client)}),
            contextlib.redirect_stdout(io.StringIO()) as messages,
        ):
            self.assertEqual(main(), 0)
        self.assertTrue(self.client.closed)
        self.assertTrue(output.exists())
        self.assertIn("Exported 1 checked rows", messages.getvalue())
        actual = load_workbook(output)
        self.assertEqual(actual.active["I2"].value, "matched")
        actual.close()

    def test_cli_rejects_in_place_output(self):
        source = self.directory / "source.xlsx"
        source.write_bytes(b"original")
        with (
            patch("export_checking_results.parse_args", return_value=argparse.Namespace(
                input_file=source, output=source, remarks_from=None,
            )),
            contextlib.redirect_stderr(io.StringIO()),
        ):
            self.assertEqual(main(), 1)
        self.assertEqual(source.read_bytes(), b"original")

    def test_remark_is_blank_manual_column_and_does_not_change_result(self):
        workbook = self.workbook()
        self.addCleanup(workbook.close)
        summary = check_workbook(workbook, self.checker)
        sheet = workbook.active
        column = len(HEADERS) + len(RESULT_COLUMNS) + 1
        self.assertEqual(sheet.cell(1, column).value, REMARK_COLUMN)
        self.assertEqual(sheet.cell(2, column).value, "")
        self.assertEqual(summary["matched"], 1)
        self.assertEqual(sheet["I2"].value, "matched")

    def test_existing_manual_remark_is_preserved_without_duplicate_column(self):
        workbook = self.workbook()
        self.addCleanup(workbook.close)
        sheet = workbook.active
        sheet["G1"] = REMARK_COLUMN
        sheet["G2"] = 'iBMS Parameters: "Battery Voltage " vs "Battery Voltage"'
        sheet["G2"].fill = PatternFill("solid", fgColor="123456")
        check_workbook(workbook, self.checker)
        self.assertEqual(sheet["G2"].value,
                         'iBMS Parameters: "Battery Voltage " vs "Battery Voltage"')
        self.assertEqual(sheet["G2"].fill.fgColor.rgb[-6:], "123456")
        self.assertEqual(sum(c.value == REMARK_COLUMN for c in sheet[1]), 1)

    def test_read_remarks_preserves_literal_whitespace_and_blank_rows(self):
        text = '  iBMS Parameters: "Battery Voltage " vs "Battery Voltage"  '
        path = self.remarks_file([
            ["APB", "APB-ACC-04-S", text],
            [None, None, None],
            [None, None, ""],
        ])
        original = path.read_bytes()
        self.assertEqual(read_remarks(path), {("APB", "APB-ACC-04-S"): text})
        self.assertEqual(path.read_bytes(), original)

    def test_read_remarks_csv_preserves_quotes_and_formula_like_literal_text(self):
        path = self.directory / "notes.csv"
        text = '=manual: "Battery Voltage " vs "Battery Voltage"'
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow([BUILDING_COLUMN, POINT_COLUMN, REMARK_COLUMN])
            writer.writerow(["APB", "APB-GEN-01-BATV", text])
        self.assertEqual(read_remarks(path),
                         {("APB", "APB-GEN-01-BATV"): text})

    def test_import_matches_building_and_full_point_not_row_order(self):
        notes = {
            ("APB", "APB-ACC-04-S"): "APB S note",
            ("OTHER", "APB-ACC-04-S"): "OTHER S note",
            ("APB", "APB-ACC-04-S1"): "Device note",
        }
        workbook = self.workbook()
        self.addCleanup(workbook.close)
        sheet = workbook.active
        sheet["A2"] = "OTHER"
        sheet.append(["APB", "ACC", "APB-ACC-04-S1"])
        sheet.append(["APB", "ACC", "APB-ACC-04-S"])
        with (
            contextlib.redirect_stdout(io.StringIO()),
            contextlib.redirect_stderr(io.StringIO()),
        ):
            check_workbook(workbook, self.checker, remarks=notes)
        self.assertEqual([sheet.cell(r, 17).value for r in (2, 3, 4)],
                         ["OTHER S note", "Device note", "APB S note"])

    def test_import_is_exact_and_reports_unmatched_old_notes(self):
        workbook = self.workbook()
        self.addCleanup(workbook.close)
        notes = {
            ("APB", "APB-ACC-04-S "): "Trailing-space key",
            ("apb", "APB-ACC-04-S"): "Different building case",
            ("APB", "APB-REMOVED-S"): "Removed point",
        }
        with (
            contextlib.redirect_stdout(io.StringIO()) as messages,
            contextlib.redirect_stderr(io.StringIO()) as warnings,
        ):
            check_workbook(workbook, self.checker, remarks=notes)
        self.assertEqual(workbook.active.cell(2, 17).value, "")
        self.assertIn("matched 0 of 3", messages.getvalue())
        self.assertEqual(warnings.getvalue().count("no matching input row"), 3)

    def test_identical_duplicate_notes_are_allowed_conflicting_notes_fail(self):
        path = self.remarks_file([
            ["APB", "APB-ACC-04-S", "Same"],
            ["APB", "APB-ACC-04-S", "Same"],
        ])
        self.assertEqual(read_remarks(path), {("APB", "APB-ACC-04-S"): "Same"})
        path = self.remarks_file([
            ["APB", "APB-ACC-04-S", "First"],
            ["APB", "APB-ACC-04-S", "Second"],
        ])
        with self.assertRaisesRegex(ValueError, "Conflicting Remarks"):
            read_remarks(path)

    def test_remarks_across_sheets_and_repeated_target_rows(self):
        path = self.remarks_file([["APB", "APB-ACC-04-S", "Shared"]])
        old = load_workbook(path)
        second = old.create_sheet("Second")
        second.append([BUILDING_COLUMN, POINT_COLUMN, REMARK_COLUMN])
        second.append(["APB", "APB-ACC-04-S", "Shared"])
        notes = old.create_sheet("Notes")
        notes.append(["Untouched"])
        old.save(path)
        old.close()
        imported = read_remarks(path)
        workbook = self.workbook()
        self.addCleanup(workbook.close)
        second = workbook.create_sheet("Second")
        second.append(HEADERS)
        second.append(["APB", "ACC", "APB-ACC-04-S"])
        workbook.active.append(["APB", "ACC", "APB-ACC-04-S"])
        with contextlib.redirect_stdout(io.StringIO()):
            check_workbook(workbook, self.checker, remarks=imported)
        self.assertEqual(workbook.active.cell(2, 17).value, "Shared")
        self.assertEqual(workbook.active.cell(3, 17).value, "Shared")
        self.assertEqual(second.cell(2, 17).value, "Shared")

    def test_import_fills_existing_blank_remark_without_changing_style(self):
        workbook = self.workbook()
        self.addCleanup(workbook.close)
        sheet = workbook.active
        sheet["G1"] = REMARK_COLUMN
        sheet["G2"].fill = PatternFill("solid", fgColor="123456")
        with contextlib.redirect_stdout(io.StringIO()):
            check_workbook(workbook, self.checker,
                           remarks={("APB", "APB-ACC-04-S"): "Imported"})
        self.assertEqual(sheet["G2"].value, "Imported")
        self.assertEqual(sheet["G2"].fill.fgColor.rgb[-6:], "123456")
        self.assertEqual(sum(c.value == REMARK_COLUMN for c in sheet[1]), 1)

    def test_existing_different_remark_cannot_be_overwritten(self):
        workbook = self.workbook()
        self.addCleanup(workbook.close)
        workbook.active["G1"] = REMARK_COLUMN
        workbook.active["G2"] = "Keep me"
        with self.assertRaisesRegex(ValueError, "conflicts with imported Remark"):
            check_workbook(workbook, self.checker,
                           remarks={("APB", "APB-ACC-04-S"): "New"})
        self.assertEqual(workbook.active["G2"].value, "Keep me")

    def test_bad_remark_headers_and_nonempty_note_without_key_fail(self):
        for headers, values, message in (
            ([BUILDING_COLUMN, POINT_COLUMN], ["APB", "P"], "No worksheet"),
            ([POINT_COLUMN, REMARK_COLUMN], ["P", "note"], "Remark requires"),
            ([BUILDING_COLUMN, POINT_COLUMN, REMARK_COLUMN, REMARK_COLUMN],
             ["APB", "P", "note", "note"], "duplicate required header"),
            ([BUILDING_COLUMN, POINT_COLUMN, REMARK_COLUMN],
             [None, "P", "note"], "non-empty string"),
        ):
            with self.subTest(headers=headers):
                path = self.remarks_file([values], headers)
                with self.assertRaisesRegex(ValueError, message):
                    read_remarks(path)

    def test_formula_and_numeric_remarks_are_not_silently_converted(self):
        for value in ('="note"', 42):
            with self.subTest(value=value):
                path = self.remarks_file([["APB", "P", value]])
                with self.assertRaisesRegex(ValueError, "literal text"):
                    read_remarks(path)

    def test_remark_key_formula_without_cache_fails_explicitly(self):
        path = self.remarks_file([['="APB"', "P", "note"]])
        with self.assertRaisesRegex(ValueError, "without a cached value"):
            read_remarks(path)

    def test_literal_formula_like_remark_stays_text_after_round_trip(self):
        workbook = self.workbook()
        self.addCleanup(workbook.close)
        with contextlib.redirect_stdout(io.StringIO()):
            check_workbook(workbook, self.checker,
                           remarks={("APB", "APB-ACC-04-S"): "=manual text"})
        output = self.directory / "literal.xlsx"
        save_workbook(workbook, output)
        actual = load_workbook(output)
        self.addCleanup(actual.close)
        self.assertEqual(actual.active.cell(2, 17).value, "=manual text")
        self.assertEqual(actual.active.cell(2, 17).data_type, "s")

    def test_duplicate_input_remark_headers_fail(self):
        workbook = self.workbook()
        self.addCleanup(workbook.close)
        workbook.active["G1"] = REMARK_COLUMN
        workbook.active["H1"] = REMARK_COLUMN
        with self.assertRaisesRegex(ValueError, "duplicate Remark"):
            check_workbook(workbook, self.checker)

    def test_cli_imports_manual_notes_and_recomputes_results(self):
        source = self.directory / "source.xlsx"
        workbook = self.workbook()
        workbook.save(source)
        workbook.close()
        old = self.workbook()
        check_workbook(old, self.checker)
        old.active["I2"] = "unmatched"
        old.active.cell(2, 17).value = (
            "iBMS Parameters: On/Off Status vs Off/On Status"
        )
        previous = self.directory / "previous.xlsx"
        old.save(previous)
        old.close()
        previous_bytes = previous.read_bytes()
        output = self.directory / "new.xlsx"
        with (
            patch("export_checking_results.parse_args", return_value=argparse.Namespace(
                input_file=source, output=output, remarks_from=previous,
            )),
            patch.dict(sys.modules, {"mongodb_connect": SimpleNamespace(DB=self.client)}),
            contextlib.redirect_stdout(io.StringIO()) as messages,
        ):
            self.assertEqual(main(), 0)
        self.assertIn("matched 1 of 1", messages.getvalue())
        self.assertEqual(previous.read_bytes(), previous_bytes)
        actual = load_workbook(output)
        self.addCleanup(actual.close)
        self.assertEqual(actual.active.cell(2, 17).value,
                         "iBMS Parameters: On/Off Status vs Off/On Status")
        self.assertEqual(actual.active["I2"].value, "matched")
        self.assertTrue(self.client.closed)

    def test_cli_remark_conflict_aborts_without_creating_output(self):
        previous = self.remarks_file([
            ["APB", "P", "First"], ["APB", "P", "Second"],
        ])
        output = self.directory / "new.xlsx"
        with (
            patch("export_checking_results.parse_args", return_value=argparse.Namespace(
                input_file=self.directory / "unused.xlsx",
                output=output, remarks_from=previous,
            )),
            contextlib.redirect_stderr(io.StringIO()) as errors,
        ):
            self.assertEqual(main(), 1)
        self.assertFalse(output.exists())
        self.assertIn("Conflicting Remarks", errors.getvalue())

    def test_cli_existing_note_conflict_does_not_save_partial_workbook(self):
        source = self.directory / "source.xlsx"
        workbook = self.workbook()
        workbook.active["G1"] = REMARK_COLUMN
        workbook.active["G2"] = "Keep original"
        workbook.save(source)
        workbook.close()
        source_bytes = source.read_bytes()
        previous = self.remarks_file([["APB", "APB-ACC-04-S", "Other note"]])
        output = self.directory / "new.xlsx"
        with (
            patch("export_checking_results.parse_args", return_value=argparse.Namespace(
                input_file=source, output=output, remarks_from=previous,
            )),
            patch.dict(sys.modules, {"mongodb_connect": SimpleNamespace(DB=self.client)}),
            contextlib.redirect_stderr(io.StringIO()) as errors,
        ):
            self.assertEqual(main(), 1)
        self.assertIn("conflicts with imported Remark", errors.getvalue())
        self.assertEqual(source.read_bytes(), source_bytes)
        self.assertFalse(output.exists())
        self.assertTrue(self.client.closed)

    def test_cli_parses_remarks_option(self):
        with patch.object(sys, "argv", [
            "export_checking_results.py", "points.csv",
            "--remarks-from", "previous.xlsx", "-o", "next.xlsx",
        ]):
            args = parse_args()
        self.assertEqual(args.input_file, Path("points.csv"))
        self.assertEqual(args.remarks_from, Path("previous.xlsx"))
        self.assertEqual(args.output, Path("next.xlsx"))


if __name__ == "__main__":
    unittest.main()
